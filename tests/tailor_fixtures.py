"""A small Word CV built in code, shaped like a real one (bold lead-ins, bullets, dates)."""

from __future__ import annotations

import io

from docx import Document

SUMMARY = (
    "Technical program manager with ten years of delivery experience across payments and "
    "platform teams, known for clear plans and calm execution."
)
COMPETENCIES = ["Program management", "Risk management", "Stakeholder management", "Agile"]
BULLETS = [
    ("Payments launch", ": Led a launch across 6 teams and cut release time by 30% in a year."),
    ("Platform", ": Ran quarterly planning for 4 squads and kept scope changes visible."),
    ("Reliability", ": Set up incident reviews that reduced repeat outages across services."),
]


def make_cv() -> bytes:
    doc = Document()
    doc.add_paragraph("Asha Rao")
    doc.add_paragraph("asha@example.com | +91 90000 00000")
    doc.add_paragraph(SUMMARY)
    doc.add_paragraph("CORE COMPETENCIES")
    doc.add_paragraph(" • ".join(COMPETENCIES))
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Acme Pay, Bengaluru\tJan 2020 - Present")
    for label, text in BULLETS:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(label).bold = True
        paragraph.add_run(text)
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech, 2012")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


LETTER_OPENING = (
    "I am excited to apply for the role at your company and bring ten years of delivery."
)
LETTER_LEAD = "A quick snapshot of what I would bring:"
LETTER_BULLETS = [
    ("Ownership", ": Led work across 6 teams from plan to launch.\n"),
    ("Planning", ": Ran quarterly planning for 4 squads.\n"),
    ("Reliability", ": Set up incident reviews.\n"),
    ("Delivery", ": Cut release time by 30% in a year.\n"),
]
LETTER_CLOSING = "Thank you for your time. I would welcome a conversation."


def make_letter() -> bytes:
    from docx.shared import Pt

    doc = Document()
    doc.add_paragraph("Asha Rao")
    doc.add_paragraph("asha@example.com | +91 90000 00000")
    doc.add_paragraph("Jan 11, 2026")
    doc.add_paragraph("Globex").runs[0].bold = True
    doc.add_paragraph("Dear Hiring Manager,")
    doc.add_paragraph(LETTER_OPENING)
    doc.add_paragraph(LETTER_LEAD)
    count = len(LETTER_BULLETS)
    for number, (label, text) in enumerate(LETTER_BULLETS):
        paragraph = doc.add_paragraph(style="List Bullet")
        if number == 0:
            paragraph.paragraph_format.space_before = Pt(12)
        if number == count - 1:
            paragraph.paragraph_format.space_after = Pt(12)
        paragraph.add_run(label + ":" if not label.endswith(":") else label).bold = True
        paragraph.add_run(text)
    doc.add_paragraph(LETTER_CLOSING)
    doc.add_paragraph("Warm regards,\nAsha Rao")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
