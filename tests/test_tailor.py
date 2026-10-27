from __future__ import annotations

import io
import json
from collections.abc import Iterator

import pytest
import requests
from docx import Document
from flask.testing import FlaskClient

from jobscan import db
from jobscan.config import Config
from jobscan.models import JobDetail
from jobscan.store import Store
from jobscan.tailor import generate, output_name, run_request, start_request
from jobscan.tailor_ai import (
    SYSTEM_PROMPT,
    TailorClient,
    TailorError,
    build_user_message,
    make_client,
    parse_proposal,
    propose,
)
from jobscan.tailor_docx import (
    DocxError,
    EditPlan,
    analyse,
    apply_plan,
    check_rewrite,
    longer_in_pages,
    verify_same_shape,
)
from jobscan.web import create_app
from jobscan.web.diffs import moved_items, word_diff
from tests.tailor_fixtures import BULLETS, COMPETENCIES, SUMMARY, make_cv
from tests.test_web import add_role

WHEN = "2026-10-16T09:00:00+05:30"


class FakeClient:
    def __init__(self, *replies: str | Exception) -> None:
        self.replies = list(replies)
        self.sent: list[str] = []

    def complete(self, system: str, user: str) -> str:
        self.sent.append(user)
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        return reply


def good_reply() -> str:
    return json.dumps(
        {
            "summary": SUMMARY.replace("calm execution", "execution"),
            "competencies_order": list(reversed(COMPETENCIES)),
            "jobs": [
                {
                    "id": "j1",
                    "order": ["j1b2", "j1b1", "j1b3"],
                    "rewrites": {"j1b3": "Set up incident reviews that cut repeat outages."},
                }
            ],
            "gaps": ["Kubernetes"],
        }
    )


# ---- Word file --------------------------------------------------------------------------------


def test_analyse_finds_the_parts() -> None:
    s = analyse(make_cv())
    assert s.summary == SUMMARY
    assert list(s.competencies) == COMPETENCIES
    assert [b.label for b in s.jobs[0].bullets] == [label for label, _ in BULLETS]
    assert all(b.editable for b in s.jobs[0].bullets)


def test_not_a_word_file_is_refused() -> None:
    with pytest.raises(DocxError):
        analyse(b"not a docx")


@pytest.mark.parametrize(
    ("new", "why"),
    [
        ("", "empty"),
        ("same", "unchanged"),
        ("a\nb", "line"),
        ("x" * 500, "longer"),
        ("short 99", "number"),
    ],
)
def test_check_rewrite_rejects(new: str, why: str) -> None:
    assert check_rewrite("same", new) is not None


def test_check_rewrite_accepts_a_shorter_text_with_known_numbers() -> None:
    assert check_rewrite("Cut time by 30% across 6 teams.", "Cut time by 30% in 6 teams.") is None


def test_apply_plan_changes_text_and_keeps_shape() -> None:
    original = make_cv()
    plan = EditPlan(
        summary=SUMMARY.replace("calm execution", "execution"),
        competencies=tuple(reversed(COMPETENCIES)),
        rewrites={"j1b3": "Set up incident reviews that cut repeat outages."},
        orders={"j1": ("j1b2", "j1b1", "j1b3")},
    )
    edited = apply_plan(original, plan)
    assert verify_same_shape(original, edited) is None
    lines = [p.text for p in Document(io.BytesIO(edited)).paragraphs]
    assert any(line.startswith("Platform") for line in lines)
    bullets = [line for line in lines if line.startswith(("Platform", "Payments", "Reliab"))]
    assert bullets[0].startswith("Platform")
    assert "Agile" in lines[4].split("•")[0]
    assert "calm" not in lines[2]


def _pdf(pages: int) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_longer_result_is_blocked_by_page_count() -> None:
    original = make_cv()
    edited = apply_plan(original, EditPlan(summary=SUMMARY.replace("calm execution", "execution")))

    def convert(data: bytes) -> bytes | None:
        return _pdf(1) if data == original else _pdf(2)

    assert longer_in_pages(original, edited, convert) is not None
    assert longer_in_pages(original, edited, lambda d: _pdf(1)) is None
    assert longer_in_pages(original, edited, lambda d: None) is None


# ---- AI ---------------------------------------------------------------------------------------


def test_prompt_has_no_contact_details() -> None:
    message = build_user_message(analyse(make_cv()), "TPM", "Acme", "Lead delivery")
    assert "asha@example.com" not in message
    assert "+91" not in message
    assert "Asha" not in message
    assert "Lead delivery" in message
    assert "Never add" in SYSTEM_PROMPT


def test_parse_keeps_good_changes_and_drops_bad_ones() -> None:
    s = analyse(make_cv())
    reply = json.dumps(
        {
            "summary": SUMMARY + " Plus a lot of extra words that make it longer than before.",
            "competencies_order": ["Agile", "Something new"],
            "jobs": [
                {"id": "j1", "order": None, "rewrites": {"j1b1": "Led 99 teams.", "zz": "x"}},
                {"id": "nope"},
            ],
            "gaps": ["a", "b", "c", "d", "e", "f", "g"],
        }
    )
    proposal = parse_proposal(s, reply)
    assert proposal.changes == ()
    assert proposal.dropped == 5
    assert len(proposal.gaps) == 6


def test_parse_good_reply() -> None:
    proposal = parse_proposal(analyse(make_cv()), "```json\n" + good_reply() + "\n```")
    assert [c["kind"] for c in proposal.changes] == [
        "summary",
        "competencies",
        "bullet",
        "order",
    ]
    assert proposal.dropped == 0
    assert type(proposal).from_json(proposal.to_json()) == proposal


def test_propose_retries_once_on_unreadable_reply() -> None:
    client = FakeClient("sorry, no json", good_reply())
    assert propose(client, analyse(make_cv()), "TPM", "Acme", "desc").changes  # type: ignore[arg-type]
    assert len(client.sent) == 2
    with pytest.raises(TailorError):
        propose(FakeClient("x", "y"), analyse(make_cv()), "TPM", "Acme", "d")  # type: ignore[arg-type]


class FakeSession:
    def __init__(self, status: int = 200, body: object = None, error: Exception | None = None):
        self.status, self.body, self.error = status, body, error

    def post(self, *args: object, **kwargs: object) -> requests.Response:
        if self.error:
            raise self.error
        response = requests.Response()
        response.status_code = self.status
        response._content = json.dumps(self.body or {}).encode()
        return response


@pytest.mark.parametrize(
    ("session", "fragment"),
    [
        (FakeSession(401), "refused the key"),
        (FakeSession(404), "not found"),
        (FakeSession(429), "rate limit"),
        (FakeSession(503), "problem"),
        (FakeSession(error=requests.Timeout()), "too long"),
        (FakeSession(error=requests.ConnectionError()), "could not be reached"),
        (FakeSession(200, {"nope": 1}), "expected form"),
    ],
)
def test_client_failures_have_plain_reasons(session: FakeSession, fragment: str) -> None:
    client = TailorClient("https://x/v1", "m", "k", session=session)  # type: ignore[arg-type]
    with pytest.raises(TailorError, match=fragment):
        client.complete("s", "u")


def test_client_returns_the_reply_text() -> None:
    body = {"choices": [{"message": {"content": "hello"}}]}
    client = TailorClient("https://x/v1", "m", "k", session=FakeSession(200, body))  # type: ignore[arg-type]
    assert client.complete("s", "u") == "hello"


def test_missing_key_is_explained() -> None:
    with pytest.raises(TailorError, match="CEREBRAS_API_KEY"):
        make_client("https://x/v1", "m", "CEREBRAS_API_KEY", {})


# ---- orchestration ----------------------------------------------------------------------------


@pytest.fixture
def store() -> Iterator[Store]:
    conn = db.connect(":memory:")
    db.migrate(conn)
    yield Store(conn)
    conn.close()


@pytest.fixture
def ready_store(store: Store) -> Store:
    run = store.start_run(window_seconds=86400, keywords_total=1)
    add_role(store, run, "1")
    add_role(store, run, "2")
    store.save_detail("LinkedIn", "1", JobDetail("Lead delivery"), when=WHEN)
    return store


def upload(store: Store) -> None:
    from jobscan.models import CvFile

    store.save_cv_file(CvFile("cv.docx", make_cv(), "sha1", WHEN))


def test_start_cases(ready_store: Store) -> None:
    store = ready_store
    assert start_request(store, "LinkedIn", "9").kind == "no_role"
    assert start_request(store, "LinkedIn", "1").kind == "no_cv"
    upload(store)
    assert start_request(store, "LinkedIn", "2").kind == "no_description"
    first = start_request(store, "LinkedIn", "1")
    assert first.kind == "started"
    again = start_request(store, "LinkedIn", "1")
    assert (again.kind, again.request_id) == ("existing", first.request_id)


def test_run_then_generate_then_new_row_after_ready(ready_store: Store) -> None:
    store = ready_store
    upload(store)
    request_id = start_request(store, "LinkedIn", "1").request_id
    assert request_id is not None
    run_request(store, request_id, Config(), lambda c: FakeClient(good_reply()))  # type: ignore[arg-type,return-value]
    row = store.get_tailor_request(request_id)
    assert row is not None
    assert row.status == "review"
    assert row.proposed_count == 4
    assert store.count_tailor_review() == 1
    assert generate(store, request_id, {"c1", "c3"}, lambda d: None) == 2
    done = store.get_tailor_request(request_id)
    assert done is not None
    assert done.status == "ready"
    assert done.applied_count == 2
    name, data = store.get_tailor_output(request_id) or ("", b"")
    assert name == "cv_Acme_Technical_Program_Manager_1.docx"
    assert data
    assert start_request(store, "LinkedIn", "1").request_id != request_id


def test_generate_needs_a_tick_and_a_review_state(ready_store: Store) -> None:
    store = ready_store
    upload(store)
    request_id = start_request(store, "LinkedIn", "1").request_id or 0
    with pytest.raises(TailorError, match="not waiting"):
        generate(store, request_id, {"c1"})
    run_request(store, request_id, Config(), lambda c: FakeClient(good_reply()))  # type: ignore[arg-type,return-value]
    with pytest.raises(TailorError, match="Tick"):
        generate(store, request_id, set(), lambda d: None)


@pytest.mark.parametrize(
    ("reply", "fragment"),
    [
        (TailorError("The AI service is busy (rate limit). Try again in a minute."), "rate limit"),
        (RuntimeError("boom"), "Unexpected"),
    ],
)
def test_failures_are_recorded_with_a_reason(
    ready_store: Store, reply: Exception, fragment: str
) -> None:
    store = ready_store
    upload(store)
    request_id = start_request(store, "LinkedIn", "1").request_id or 0
    run_request(store, request_id, Config(), lambda c: FakeClient(reply))  # type: ignore[arg-type,return-value]
    row = store.get_tailor_request(request_id)
    assert row is not None
    assert row.status == "failed"
    assert fragment in row.note


def test_replacing_the_cv_makes_waiting_requests_stale(ready_store: Store) -> None:
    store = ready_store
    upload(store)
    request_id = start_request(store, "LinkedIn", "1").request_id or 0
    run_request(store, request_id, Config(), lambda c: FakeClient(good_reply()))  # type: ignore[arg-type,return-value]
    store.mark_tailor_stale()
    row = store.get_tailor_request(request_id)
    assert row is not None
    assert row.status == "stale"
    assert store.count_tailor_review() == 0


def test_output_name_is_safe() -> None:
    assert output_name("My CV.docx", "A/B Co.", "Sr. TPM (Pune)") == "My_CV_A_B_Co_Sr_TPM_Pune.docx"


def test_word_diff_marks_changes() -> None:
    before, after = word_diff("led big teams", "led teams")
    assert "<del>big </del>" in before or "<del>big</del>" in before
    assert "<ins>" not in after
    assert moved_items(["a", "b"], ["b", "a"]) == [("b", True), ("a", True)]


# ---- web --------------------------------------------------------------------------------------


@pytest.fixture
def client(ready_store: Store) -> FlaskClient:
    app = create_app(
        ready_store,
        secret_key="t",
        make_tailor_client=lambda c: FakeClient(good_reply()),  # type: ignore[arg-type,return-value]
        tailor_convert=lambda d: None,
    )
    return app.test_client()


def post_cv(client: FlaskClient, name: str, data: bytes) -> str:
    response = client.post(
        "/cv",
        data={"cv": (io.BytesIO(data), name)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    return response.get_data(as_text=True)


def test_pdf_or_text_cv_cannot_be_tailored(client: FlaskClient) -> None:
    page = post_cv(
        client, "cv.txt", b"Worked at Acme\nLed a launch of a product across teams.\n" * 3
    )
    assert "Not available" in page
    response = client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    assert response.headers["Location"].endswith("/cv")


def test_whole_flow_through_the_web(client: FlaskClient) -> None:
    assert "Available" in post_cv(client, "cv.docx", make_cv())
    response = client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    assert response.headers["Location"].endswith("/cv/history#r1")
    history = client.get("/cv/history").get_data(as_text=True)
    assert "Needs your review" in history
    assert 'class="badge"' in history
    review = client.get("/cv/tailor/1").get_data(as_text=True)
    assert review.count('name="change"') == 4
    assert "Kubernetes" in review
    # a second tap goes to the same row
    again = client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    assert again.headers["Location"].endswith("#r1")
    done = client.post(
        "/cv/tailor/1/generate", data={"change": ["c1", "c2"]}, follow_redirects=True
    )
    assert "Resume ready with 2 changes" in done.get_data(as_text=True)
    download = client.get("/cv/tailor/1/download")
    assert download.status_code == 200
    assert "attachment" in download.headers["Content-Disposition"]
    assert Document(io.BytesIO(download.data)).paragraphs
    assert "Ready" in client.get("/cv/history").get_data(as_text=True)


def test_generate_without_ticks_stays_on_review(client: FlaskClient) -> None:
    post_cv(client, "cv.docx", make_cv())
    client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    page = client.post("/cv/tailor/1/generate", data={}, follow_redirects=True)
    assert "Tick at least one change" in page.get_data(as_text=True)


def test_discard_then_start_again_makes_a_new_row(client: FlaskClient) -> None:
    post_cv(client, "cv.docx", make_cv())
    client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    client.post("/cv/tailor/1/discard")
    assert "Discarded" in client.get("/cv/history").get_data(as_text=True)
    response = client.post("/cv/tailor/1/again")
    assert response.headers["Location"].endswith("#r2")


def test_new_cv_makes_the_review_stale(client: FlaskClient) -> None:
    post_cv(client, "cv.docx", make_cv())
    client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "1"})
    post_cv(client, "cv.txt", b"Different CV text with lots of words about delivery work.\n" * 4)
    assert "CV changed, start again" in client.get("/cv/history").get_data(as_text=True)


def test_panel_button_and_missing_description(client: FlaskClient) -> None:
    with_text = client.get("/roles/LinkedIn/1").get_data(as_text=True)
    assert "Tailor resume" in with_text
    assert "Needs a job description" not in with_text
    without = client.get("/roles/LinkedIn/2").get_data(as_text=True)
    assert "Needs a job description first" in without


def test_unknown_requests_are_404(client: FlaskClient) -> None:
    assert client.get("/cv/tailor/99").status_code == 404
    assert client.get("/cv/tailor/99/download").status_code == 404
    assert (
        client.post("/tailor/start", data={"portal": "LinkedIn", "job_id": "zz"}).status_code == 404
    )


def test_unreadable_word_file_shows_its_reason(
    client: FlaskClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    import jobscan.web.app as web

    def broken(data: bytes) -> None:
        raise DocxError("The python-docx package is missing: pip install python-docx")

    monkeypatch.setattr(web, "analyse", broken)
    page = post_cv(client, "cv.docx", make_cv())
    assert "python-docx package is missing" in page
    start = client.post(
        "/tailor/start", data={"portal": "LinkedIn", "job_id": "1"}, follow_redirects=True
    )
    assert "cannot be tailored" in start.get_data(as_text=True)
