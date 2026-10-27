"""Cover letters: the Word engine, the AI answer checks, the request flow and the web pages."""

from __future__ import annotations

import io
import json
from collections.abc import Iterator
from datetime import date

import pytest
from docx import Document
from flask.testing import FlaskClient

from jobscan import db
from jobscan.config import Config
from jobscan.cover_ai import parse_cover, propose_cover
from jobscan.cover_docx import (
    LetterBullet,
    LetterPlan,
    analyse_letter,
    build_letter,
    check_text,
    format_date,
    verify_letter,
)
from jobscan.models import CvFile, JobDetail
from jobscan.store import Store
from jobscan.tailor import (
    CoverEdit,
    generate_cover,
    run_request,
    start_request,
)
from jobscan.tailor_ai import TailorError
from jobscan.tailor_docx import DocxError, analyse
from jobscan.web.app import create_app
from tests.tailor_fixtures import LETTER_BULLETS, make_cv, make_letter
from tests.test_tailor import FakeClient
from tests.test_web import add_role

WHEN = "2026-10-16T09:00:00+05:30"


def reply(bullets: int = 5, **over: object) -> str:
    data: dict[str, object] = {
        "opening": "I would love to lead delivery at Acme with ten years behind me.",
        "lead_in": "What I would bring:",
        "bullets": [
            {
                "label": f"Point {chr(97 + i)}",
                "text": f"Led work across 6 teams, item {chr(97 + i)}.",
            }
            for i in range(bullets)
        ],
        "closing": "Thank you. I would welcome a conversation.",
        "gaps": ["Kubernetes"],
    }
    data.update(over)
    return json.dumps(data)


# ---- Word engine ------------------------------------------------------------------------------


def test_analyse_letter_finds_the_parts() -> None:
    s = analyse_letter(make_letter())
    assert s.date_text == "Jan 11, 2026"
    assert s.company_text == "Globex"
    assert len(s.bullets) == len(LETTER_BULLETS)
    assert s.bullets[0].label == "Ownership"
    assert s.lead_in.startswith("A quick snapshot")


def test_letter_without_greeting_is_refused() -> None:
    doc = Document()
    doc.add_paragraph("Hello")
    out = io.BytesIO()
    doc.save(out)
    with pytest.raises(DocxError):
        analyse_letter(out.getvalue())


@pytest.mark.parametrize("count", [1, 2, 4, 7])
def test_build_letter_any_bullet_count_keeps_layout(count: int) -> None:
    data = make_letter()
    plan = LetterPlan(
        "Open.",
        "Lead:",
        tuple(LetterBullet(f"L{i}", f"T{i}.") for i in range(count)),
        "Close.",
        format_date(date(2026, 3, 5)),
        "Acme",
    )
    out = build_letter(data, plan)
    assert verify_letter(data, out) is None
    s = analyse_letter(out)
    assert len(s.bullets) == count
    assert s.date_text == "Mar 5, 2026"
    assert s.company_text == "Acme"
    assert s.opening == "Open."
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    assert texts[0] == "Asha Rao"
    assert "Dear Hiring Manager," in texts
    assert texts[-1].startswith("Warm regards")
    assert analyse_letter(data).bullets[0].label == "Ownership"  # the input is untouched


def test_check_text_rules() -> None:
    allowed = {"6", "30"}
    assert check_text("text", "Led 6 teams.", allowed) is None
    assert check_text("text", "Led 9 teams.", allowed)
    assert check_text("text", "Two\nlines", allowed)
    assert check_text("text", "**bold**", allowed)
    assert check_text("lead", "x" * 500, allowed)
    assert check_text("text", "  ", allowed)


# ---- AI answer --------------------------------------------------------------------------------


def parts() -> tuple[object, object, str]:
    return analyse(make_cv()), analyse_letter(make_letter()), "ten years 6 teams 4 30%"


def test_parse_cover_good_and_bad_parts() -> None:
    template = analyse_letter(make_letter())
    ok = parse_cover(template, "ten years 6 teams 4 30%", reply())
    assert [c["kind"] for c in ok.changes].count("bullet") == 5
    assert ok.gaps == ("Kubernetes",)
    bad = json.loads(reply())
    bad["bullets"][0]["text"] = "Grew revenue by 99%."
    out = parse_cover(template, "ten years 6 teams 4 30%", json.dumps(bad))
    assert [c["kind"] for c in out.changes].count("bullet") == 4
    assert out.dropped == 1


def test_parse_cover_needs_enough_bullets() -> None:
    template = analyse_letter(make_letter())
    with pytest.raises(ValueError, match="usable bullets"):
        parse_cover(template, "6", reply(bullets=2))


def test_propose_cover_retries_once_then_fails() -> None:
    structure, template = analyse(make_cv()), analyse_letter(make_letter())
    client = FakeClient("nonsense", reply())
    assert propose_cover(client, structure, template, "PM", "Acme", "desc")  # type: ignore[arg-type]
    assert len(client.sent) == 2
    with pytest.raises(TailorError):
        propose_cover(FakeClient("a", "b"), structure, template, "PM", "Acme", "d")  # type: ignore[arg-type]


def test_prompt_has_no_contact_details() -> None:
    structure, template = analyse(make_cv()), analyse_letter(make_letter())
    client = FakeClient(reply())
    propose_cover(client, structure, template, "PM", "Acme", "desc")  # type: ignore[arg-type]
    sent = client.sent[0]
    assert "asha@example.com" not in sent
    assert "90000" not in sent
    assert "Asha" not in sent
    assert "Dear Hiring Manager" not in sent


# ---- requests ---------------------------------------------------------------------------------


@pytest.fixture
def store() -> Iterator[Store]:
    conn = db.connect(":memory:")
    db.migrate(conn)
    s = Store(conn)
    run = s.start_run(window_seconds=86400, keywords_total=1)
    add_role(s, run, "1")
    add_role(s, run, "2")
    s.save_detail("LinkedIn", "1", JobDetail("Lead delivery"), when=WHEN)
    yield s
    conn.close()


def upload_both(store: Store, template: bytes | None = None) -> None:
    store.save_cv_file(CvFile("cv.docx", make_cv(), "cvsha", WHEN))
    store.save_cover_template(
        CvFile("Asha_Cover_Letter.docx", template or make_letter(), "tsha", WHEN)
    )


def test_start_cases(store: Store) -> None:
    assert start_request(store, "LinkedIn", "1", "cover").kind == "no_cv"
    store.save_cv_file(CvFile("cv.docx", make_cv(), "cvsha", WHEN))
    assert start_request(store, "LinkedIn", "1", "cover").kind == "no_template"
    store.save_cover_template(CvFile("t.docx", make_letter(), "tsha", WHEN))
    assert start_request(store, "LinkedIn", "2", "cover").kind == "no_description"
    first = start_request(store, "LinkedIn", "1", "cover")
    assert first.kind == "started"
    assert start_request(store, "LinkedIn", "1", "cover").request_id == first.request_id
    # a resume request for the same role is a separate row
    resume = start_request(store, "LinkedIn", "1")
    assert resume.kind == "started"
    assert resume.request_id != first.request_id


def prepared(store: Store) -> int:
    upload_both(store)
    request_id = start_request(store, "LinkedIn", "1", "cover").request_id
    assert request_id is not None
    run_request(store, request_id, Config(), lambda c: FakeClient(reply(6)))  # type: ignore[arg-type,return-value]
    return request_id


def edits_for(store: Store, request_id: int, **changes: CoverEdit) -> dict[str, CoverEdit]:
    return dict(changes)


def test_run_and_generate(store: Store) -> None:
    request_id = prepared(store)
    row = store.get_tailor_request(request_id)
    assert row is not None
    assert (row.kind, row.status) == ("cover", "review")
    count = generate_cover(store, request_id, {}, lambda d: None, date(2026, 10, 27))
    assert count == 6
    name, data = store.get_tailor_output(request_id) or ("", b"")
    assert name == "Asha_Cover_Letter_Acme_Technical_Program_Manager_1.docx"
    s = analyse_letter(data)
    assert s.date_text == "Oct 27, 2026"
    assert s.company_text == "Acme"
    assert len(s.bullets) == 6
    assert start_request(store, "LinkedIn", "1", "cover").request_id != request_id


def test_unticked_bullets_are_left_out_and_edits_are_used(store: Store) -> None:
    request_id = prepared(store)
    proposal = json.loads(store.get_tailor_request(request_id).proposal)  # type: ignore[union-attr]
    ids = [c["id"] for c in proposal["changes"] if c["kind"] == "bullet"]
    opening = next(c["id"] for c in proposal["changes"] if c["kind"] == "opening")
    edits = {
        ids[0]: CoverEdit("Point 0", "x", on=False),
        opening: CoverEdit("", "A hand written opening.", True),
    }
    assert generate_cover(store, request_id, edits, lambda d: None) == 5
    _, data = store.get_tailor_output(request_id) or ("", b"")
    assert analyse_letter(data).opening == "A hand written opening."


def test_bad_edit_is_refused_and_nothing_changes(store: Store) -> None:
    request_id = prepared(store)
    proposal = json.loads(store.get_tailor_request(request_id).proposal)  # type: ignore[union-attr]
    closing = next(c["id"] for c in proposal["changes"] if c["kind"] == "closing")
    with pytest.raises(TailorError, match="number"):
        generate_cover(
            store, request_id, {closing: CoverEdit("", "We grew 77% fast.")}, lambda d: None
        )
    row = store.get_tailor_request(request_id)
    assert row is not None
    assert row.status == "review"


def test_too_many_pages_is_blocked(store: Store) -> None:
    import io as _io

    from pypdf import PdfWriter

    request_id = prepared(store)
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(100, 100)
    buf = _io.BytesIO()
    writer.write(buf)
    with pytest.raises(TailorError, match="3 pages"):
        generate_cover(store, request_id, {}, lambda d: buf.getvalue())


def test_replacing_template_stales_only_cover_requests(store: Store) -> None:
    request_id = prepared(store)
    resume_id = start_request(store, "LinkedIn", "1").request_id
    assert resume_id is not None
    store.mark_tailor_stale("cover", "Your cover letter template was replaced. Start again.")
    cover = store.get_tailor_request(request_id)
    resume = store.get_tailor_request(resume_id)
    assert cover is not None
    assert resume is not None
    assert cover.status == "stale"
    assert resume.status == "preparing"
    store.mark_tailor_stale()
    assert store.get_tailor_request(resume_id).status == "stale"  # type: ignore[union-attr]


def test_template_replaced_while_waiting_fails_generate(store: Store) -> None:
    request_id = prepared(store)
    store.save_cover_template(CvFile("new.docx", make_letter(), "othersha", WHEN))
    with pytest.raises(TailorError, match="template was replaced"):
        generate_cover(store, request_id, {}, lambda d: None)


# ---- web --------------------------------------------------------------------------------------


@pytest.fixture
def client(store: Store) -> FlaskClient:
    app = create_app(
        store,
        secret_key="t",
        make_tailor_client=lambda c: FakeClient(reply(6)),  # type: ignore[arg-type,return-value]
        tailor_convert=lambda d: None,
    )
    return app.test_client()


def post_file(client: FlaskClient, url: str, field: str, name: str, data: bytes) -> str:
    response = client.post(
        url,
        data={field: (io.BytesIO(data), name)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    return response.get_data(as_text=True)


def test_missing_template_sends_you_to_upload(client: FlaskClient) -> None:
    post_file(client, "/cv", "cv", "cv.docx", make_cv())
    response = client.post(
        "/tailor/start", data={"portal": "LinkedIn", "job_id": "1", "kind": "cover"}
    )
    assert response.headers["Location"].endswith("/cv")
    page = client.get("/cv").get_data(as_text=True)
    assert "Upload your cover letter template (Word) to tailor a cover letter." in page


def test_template_upload_checks_the_file(client: FlaskClient) -> None:
    page = post_file(client, "/cv/template", "template", "t.txt", b"hello")
    assert "must be a Word" in page
    doc = Document()
    doc.add_paragraph("No greeting here")
    buf = io.BytesIO()
    doc.save(buf)
    assert "cannot be used as a template" in post_file(
        client, "/cv/template", "template", "t.docx", buf.getvalue()
    )


def test_whole_flow_through_the_web(client: FlaskClient) -> None:
    post_file(client, "/cv", "cv", "cv.docx", make_cv())
    assert "Cover letter template saved" in post_file(
        client, "/cv/template", "template", "Letter.docx", make_letter()
    )
    response = client.post(
        "/tailor/start", data={"portal": "LinkedIn", "job_id": "1", "kind": "cover"}
    )
    assert response.headers["Location"].endswith("/cv/history#r1")
    history = client.get("/cv/history").get_data(as_text=True)
    assert "Cover letter" in history
    assert "Needs your review" in history
    review = client.get("/cv/tailor/1").get_data(as_text=True)
    assert review.count('name="keep_') == 6
    assert "Kubernetes" in review
    assert 'with <span id="n">6</span> bullets' in review
    # a bad edit keeps the person's text and says why
    bad = client.post("/cv/tailor/1/generate", data={"text_c1": "We grew 77% overnight."})
    assert bad.status_code == 200
    text = bad.get_data(as_text=True)
    assert "We grew 77% overnight." in text
    assert "not accepted" in text
    done = client.post(
        "/cv/tailor/1/generate",
        data={"keep_c3": "on", "keep_c4": "on", "keep_c5": "on"},
        follow_redirects=True,
    )
    assert "Cover letter ready with 3 bullets" in done.get_data(as_text=True)
    download = client.get("/cv/tailor/1/download")
    assert download.status_code == 200
    assert len(analyse_letter(download.data).bullets) == 3
    assert "View letter" in client.get("/cv/history").get_data(as_text=True)


def test_panel_has_both_buttons(client: FlaskClient) -> None:
    page = client.get("/roles/LinkedIn/1").get_data(as_text=True)
    assert "Tailor resume" in page
    assert "Tailor cover letter" in page
