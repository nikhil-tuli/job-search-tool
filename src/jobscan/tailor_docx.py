"""Read a Word CV's structure and edit it in place, keeping every bit of its formatting.

Only the text inside existing runs changes, and bullets are moved, never rebuilt. So fonts,
bullets, bold labels, spacing and page layout stay exactly as the person made them.

What can change: the summary wording, the order of the core competencies, the wording of a
bullet after its bold label, and the order of bullets inside one job. Nothing else is touched.
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

MAX_DOCX_PARAGRAPHS = 2_000
_LEAD = re.compile(r"^[\s.:;,\-\u2013\u2014]*")
_TAIL = re.compile(r"\s*$")
_NUMBER = re.compile(r"\d[\d,.]*\d|\d")
_SEPARATOR = re.compile(r"(\s*•\s*)")
_DATE = re.compile(r"(?:19|20)\d\d|['\u2019]\d\d")
_COMPETENCY_HEADING = re.compile(
    r"^(?:core |key |technical )?(?:competencies|skills)\b|^key skills\b", re.I
)
_EXPERIENCE_HEADING = re.compile(
    r"^(?:professional |work |relevant )?(?:experience|employment)\b", re.I
)
_END_HEADING = re.compile(
    r"^(?:education|certifications?|projects|awards|languages|publications|skills|"
    r"additional|achievements)\b",
    re.I,
)


class DocxError(ValueError):
    """The Word file cannot be used. The message says why, in plain words."""


@dataclass(frozen=True)
class Bullet:
    id: str
    label: str  # the bold lead-in, for display only
    lead: str  # punctuation before the text, kept as it is (for example ": ")
    core: str  # the text the AI may reword
    tail: str  # trailing space, kept as it is
    editable: bool


@dataclass(frozen=True)
class Job:
    id: str
    header: str  # employer and place, without the dates
    bullets: tuple[Bullet, ...]


@dataclass(frozen=True)
class CvStructure:
    summary: str  # the part of the summary that can be reworded ("" if none was found)
    summary_context: str  # the whole first paragraph, for the AI to read
    competencies: tuple[str, ...]
    jobs: tuple[Job, ...]
    skipped: tuple[str, ...] = field(default=())  # parts that could not be found


@dataclass(frozen=True)
class EditPlan:
    summary: str | None = None
    competencies: tuple[str, ...] | None = None
    rewrites: dict[str, str] = field(default_factory=dict)  # bullet id -> new text
    orders: dict[str, tuple[str, ...]] = field(default_factory=dict)  # job id -> bullet ids


# ---- reading --------------------------------------------------------------------------------


def _open(data: bytes) -> Any:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocxError("The python-docx package is missing: pip install python-docx") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocxError("This file could not be read as a Word document.") from exc
    if len(document.paragraphs) > MAX_DOCX_PARAGRAPHS:
        raise DocxError("This Word file is too long to be a CV.")
    return document


def _is_bullet(paragraph: Any) -> bool:
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        return True
    style = paragraph.style.name if paragraph.style is not None else ""
    return style.lower().startswith("list")


def _runs_with_text(paragraph: Any) -> list[Any]:
    return [r for r in paragraph.runs if r.text.strip()]


def _plain_runs(paragraph: Any) -> list[Any]:
    return [r for r in _runs_with_text(paragraph) if not r.bold]


def _split_ends(text: str) -> tuple[str, str, str]:
    lead = _LEAD.match(text).group(0)  # type: ignore[union-attr]
    rest = text[len(lead) :]
    tail = _TAIL.search(rest).group(0)  # type: ignore[union-attr]
    core = rest[: len(rest) - len(tail)] if tail else rest
    return lead, core, tail


def _heading(paragraph: Any, pattern: re.Pattern[str]) -> bool:
    text = paragraph.text.strip()
    return 0 < len(text) <= 40 and bool(pattern.match(text))


def _find(paragraphs: Sequence[Any], pattern: re.Pattern[str], start: int = 0) -> int | None:
    for index in range(start, len(paragraphs)):
        if _heading(paragraphs[index], pattern):
            return index
    return None


def _header_text(paragraph: Any) -> str:
    """The employer and place from a header line, without the dates after the tab."""
    return re.sub(r"\s+", " ", paragraph.text.split("\t")[0]).strip()


def _read_parts(document: Any) -> tuple[CvStructure, dict[str, Any]]:
    """The structure, and for each editable thing the paragraph or run it lives in."""
    paragraphs = list(document.paragraphs)
    handles: dict[str, Any] = {}
    skipped: list[str] = []

    comp_heading = _find(paragraphs, _COMPETENCY_HEADING)
    exp_heading = _find(paragraphs, _EXPERIENCE_HEADING)
    if exp_heading is None:
        raise DocxError(
            "Could not find a Professional Experience heading in this Word file, so it cannot "
            "be tailored."
        )

    summary, summary_context = "", ""
    limit = comp_heading if comp_heading is not None else exp_heading
    for index in range(limit):
        paragraph = paragraphs[index]
        if len(paragraph.text.strip()) >= 120 and not _is_bullet(paragraph):
            plain = _plain_runs(paragraph)
            if plain:
                run = max(plain, key=lambda r: len(r.text))
                lead, core, _tail = _split_ends(run.text)
                if len(core) >= 60:
                    summary, summary_context = core, paragraph.text.strip()
                    handles["summary"] = (paragraph, run)
            break
    if not summary:
        skipped.append("summary")

    competencies: tuple[str, ...] = ()
    if comp_heading is not None:
        for index in range(comp_heading + 1, exp_heading):
            paragraph = paragraphs[index]
            if "•" in paragraph.text:
                runs = _runs_with_text(paragraph)
                if len(runs) == 1:
                    parts = _SEPARATOR.split(runs[0].text)
                    items = tuple(p.strip() for p in parts[0::2] if p.strip())
                    if len(items) >= 2:
                        competencies = items
                        handles["competencies"] = (paragraph, runs[0])
                break
    if not competencies:
        skipped.append("core competencies")

    end = _find(paragraphs, _END_HEADING, exp_heading + 1) or len(paragraphs)
    jobs: list[Job] = []
    current_header = ""
    current: list[tuple[Any, Bullet]] = []

    def close() -> None:
        if current_header and current:
            job_id = f"j{len(jobs) + 1}"
            bullets = []
            for number, (paragraph, bullet) in enumerate(current, start=1):
                bullet_id = f"{job_id}b{number}"
                bullets.append(
                    Bullet(
                        bullet_id,
                        bullet.label,
                        bullet.lead,
                        bullet.core,
                        bullet.tail,
                        bullet.editable,
                    )
                )
                handles[bullet_id] = paragraph
            jobs.append(Job(job_id, current_header, tuple(bullets)))
            handles[job_id] = [paragraph for paragraph, _ in current]

    for index in range(exp_heading + 1, end):
        paragraph = paragraphs[index]
        if _is_bullet(paragraph):
            if not current_header:
                continue
            plain = _plain_runs(paragraph)
            editable = len(plain) == 1
            label = ""
            lead = core = tail = ""
            if editable:
                run = plain[0]
                lead, core, tail = _split_ends(run.text)
                label = "".join(r.text for r in paragraph.runs if r.bold and r.text.strip()).strip()
                editable = len(core) >= 20
            current.append((paragraph, Bullet("", label, lead, core, tail, editable)))
        elif "\t" in paragraph.text and _DATE.search(paragraph.text):
            close()
            current_header, current = _header_text(paragraph), []
    close()

    if not jobs:
        raise DocxError(
            "Could not find bulleted experience under dated employer lines in this Word file, "
            "so it cannot be tailored."
        )
    structure = CvStructure(summary, summary_context, competencies, tuple(jobs), tuple(skipped))
    # The bullet runs are looked up again when editing, from the paragraph handles.
    return structure, handles


def analyse(data: bytes) -> CvStructure:
    return _read_parts(_open(data))[0]


# ---- checks on proposed text ----------------------------------------------------------------


def numbers_in(text: str) -> set[str]:
    return {m.group(0).strip(".,") for m in _NUMBER.finditer(text)}


def check_rewrite(original: str, new: str) -> str | None:
    """Why a proposed rewording must be rejected, or None when it is acceptable."""
    if not new.strip():
        return "empty"
    if new == original:
        return "unchanged"
    if "\n" in new or "•" in new or "**" in new:
        return "has line breaks or formatting marks"
    if len(new) > len(original):
        return "longer than the original"
    if not numbers_in(new) <= numbers_in(original):
        return "has a number that is not in the original"
    return None


# ---- editing --------------------------------------------------------------------------------


def apply_plan(data: bytes, plan: EditPlan) -> bytes:
    """A copy of the Word file with the plan applied. The original bytes are never changed."""
    document = _open(data)
    structure, handles = _read_parts(document)

    if plan.summary is not None and "summary" in handles:
        _paragraph, run = handles["summary"]
        lead, _core, tail = _split_ends(run.text)
        run.text = lead + plan.summary + tail

    if plan.competencies is not None and "competencies" in handles:
        _paragraph, run = handles["competencies"]
        parts = _SEPARATOR.split(run.text)
        items = parts[0::2]
        stripped = [i.strip() for i in items]
        positions = [n for n, text in enumerate(stripped) if text]
        if sorted(plan.competencies) != sorted(stripped[n] for n in positions):
            raise DocxError("The new competency list does not match the current one.")
        for position, item in zip(positions, plan.competencies, strict=True):
            keep_lead = items[position][: len(items[position]) - len(items[position].lstrip())]
            keep_tail = items[position][len(items[position].rstrip()) :]
            items[position] = keep_lead + item + keep_tail
        parts[0::2] = items
        run.text = "".join(parts)

    bullets = {b.id: b for job in structure.jobs for b in job.bullets}
    for bullet_id, new_text in plan.rewrites.items():
        bullet = bullets.get(bullet_id)
        if bullet is None or not bullet.editable:
            raise DocxError(f"Bullet {bullet_id} cannot be edited.")
        paragraph = handles[bullet_id]
        run = _plain_runs(paragraph)[0]
        run.text = bullet.lead + new_text + bullet.tail

    for job in structure.jobs:
        order = plan.orders.get(job.id)
        if order is None:
            continue
        ids = [b.id for b in job.bullets]
        if sorted(order) != sorted(ids):
            raise DocxError(f"The new bullet order for {job.id} does not match the current one.")
        if list(order) != ids:
            _reorder([handles[i]._p for i in ids], [handles[i]._p for i in order])

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _reorder(originals: list[Any], new_order: list[Any]) -> None:
    """Put the paragraphs of `new_order` into the places the `originals` occupy."""
    from lxml import etree

    markers = []
    for element in originals:
        marker = etree.Element("marker")
        element.addprevious(marker)
        markers.append(marker)
    for element in originals:
        element.getparent().remove(element)
    for marker, element in zip(markers, new_order, strict=True):
        marker.addprevious(element)
        marker.getparent().remove(marker)


# ---- checking the result --------------------------------------------------------------------


def _shape(data: bytes) -> tuple[int, int, list[str]]:
    document = _open(data)
    chars = sum(len(p.text) for p in document.paragraphs)
    props = sorted((p._p.pPr.xml if p._p.pPr is not None else "") for p in document.paragraphs)
    return len(document.paragraphs), chars, props


def verify_same_shape(original: bytes, edited: bytes) -> str | None:
    """Why the edited file differs in layout from the original, or None when it matches."""
    before, after = _shape(original), _shape(edited)
    if before[0] != after[0]:
        return "the number of paragraphs changed"
    if after[1] > before[1]:
        return "the file became longer than the original"
    if before[2] != after[2]:
        return "paragraph formatting changed"
    return None


Converter = Callable[[bytes], bytes | None]


def libreoffice_pdf(data: bytes) -> bytes | None:
    """PDF bytes of a Word file made with LibreOffice, or None when it is not installed."""
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if binary is None:
        return None
    with tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / "cv.docx"
        source.write_bytes(data)
        try:
            subprocess.run(
                [binary, "--headless", "--convert-to", "pdf", "--outdir", folder, str(source)],
                check=True,
                capture_output=True,
                timeout=90,
            )
        except (subprocess.SubprocessError, OSError):
            return None
        target = Path(folder) / "cv.pdf"
        return target.read_bytes() if target.exists() else None


def page_count(pdf: bytes) -> int | None:
    from pypdf import PdfReader

    try:
        return len(PdfReader(io.BytesIO(pdf)).pages)
    except Exception:
        return None


def longer_in_pages(
    original: bytes, edited: bytes, convert: Converter = libreoffice_pdf
) -> str | None:
    """A message when the edited file runs to more pages than the original, else None.

    Needs LibreOffice. Without it nothing is checked and None is returned.
    """
    first, second = convert(original), convert(edited)
    if first is None or second is None:
        return None
    pages_before, pages_after = page_count(first), page_count(second)
    if pages_before is None or pages_after is None:
        return None
    if pages_after > pages_before:
        return f"the resume would run to {pages_after} pages instead of {pages_before}"
    return None
